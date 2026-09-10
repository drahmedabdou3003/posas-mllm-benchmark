"""
Stage 5: Export Results

Reads validated results and generates:
  - outputs/results_master.xlsx  (2 sheets: Main Results, Audit Log)
  - outputs/results_master.csv   (Sheet 1 only)

The model list is read dynamically from settings.yaml, so researchers can
add or remove models without editing this file. Set `display_name` on each
model in settings.yaml to control the column prefix, otherwise the key is used.
"""

import csv
import json
import logging
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.stage1_ingest import read_manifest

logger = logging.getLogger("posas")

OBSERVER_ITEMS = [
    "vascularity", "pigmentation", "thickness", "relief", "surface_area",
]

# Two blinded human raters (H1, H2). Increase / decrease this list to match
# your study protocol. Cells are left blank for researcher-provided scores.
HUMAN_RATERS = ["h1", "h2"]

HEADER_FILL = PatternFill("solid", fgColor="2F4F8F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _load_validated_results(validated_dir: str) -> dict:
    results = {}
    vdir = Path(validated_dir)
    if not vdir.exists():
        return results
    for model_dir in sorted(vdir.iterdir()):
        if not model_dir.is_dir():
            continue
        model_name = model_dir.name
        for vfile in sorted(model_dir.glob("*.json")):
            image_id = vfile.stem
            data = json.loads(vfile.read_text(encoding="utf-8"))
            results[(model_name, image_id)] = data
    return results


def _load_run_log(run_log_path: str) -> list:
    path = Path(run_log_path)
    if not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _model_display_names(settings: dict) -> dict:
    """
    Return {model_key: display_name}. Uses settings.models[k].display_name if set,
    otherwise the model key itself.
    """
    return {
        k: cfg.get("display_name", k)
        for k, cfg in settings["models"].items()
    }


def _build_sheet1_headers(model_names: list, display_names: dict) -> list:
    headers = [
        "image_id", "filename", "file_hash_sha256",
        "scar_type_label", "scar_location", "fitzpatrick_type",
    ]
    for model in model_names:
        d = display_names.get(model, model)
        for item in OBSERVER_ITEMS:
            headers.append(f"{d}_obs_{item}")
        headers.append(f"{d}_obs_total")
        headers.append(f"{d}_valid_json")
    for rater in HUMAN_RATERS:
        for item in OBSERVER_ITEMS:
            headers.append(f"{rater}_obs_{item}")
        headers.append(f"{rater}_obs_total")
    headers.append("any_plausibility_flag")
    headers.append("plausibility_notes")
    return headers


def _build_sheet1_row(manifest_row: dict, results: dict,
                      model_names: list, display_names: dict) -> dict:
    image_id = manifest_row["image_id"]
    row = {
        "image_id": image_id,
        "filename": manifest_row["original_filename"],
        "file_hash_sha256": manifest_row["file_hash_sha256"],
        "scar_type_label": manifest_row.get("scar_type_label", ""),
        "scar_location": manifest_row.get("scar_location", ""),
        "fitzpatrick_type": manifest_row.get("fitzpatrick_type", ""),
    }

    plausibility_notes = []

    for model in model_names:
        d = display_names.get(model, model)
        key = (model, image_id)
        result = results.get(key)

        if result and result.get("validation_passed") and result.get("parsed_scores"):
            scores = result["parsed_scores"]
            item_scores = []
            for item in OBSERVER_ITEMS:
                s = scores.get(item, "")
                row[f"{d}_obs_{item}"] = s
                if isinstance(s, int):
                    item_scores.append(s)
            row[f"{d}_obs_total"] = sum(item_scores) if item_scores else ""
            row[f"{d}_valid_json"] = "TRUE"
            flags = result.get("plausibility_flags", [])
            if flags:
                plausibility_notes.extend([f"{d}: {f}" for f in flags])
        else:
            for item in OBSERVER_ITEMS:
                row[f"{d}_obs_{item}"] = ""
            row[f"{d}_obs_total"] = ""
            row[f"{d}_valid_json"] = "FALSE" if result else "MISSING"

    for rater in HUMAN_RATERS:
        for item in OBSERVER_ITEMS:
            row[f"{rater}_obs_{item}"] = ""
        row[f"{rater}_obs_total"] = ""

    row["any_plausibility_flag"] = "TRUE" if plausibility_notes else "FALSE"
    row["plausibility_notes"] = " | ".join(plausibility_notes)

    return row


def _apply_header_style(ws, row_num: int, fill: PatternFill) -> None:
    for cell in ws[row_num]:
        if cell.value is not None:
            cell.font = HEADER_FONT
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(ws, min_width: int = 8, max_width: int = 40) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val_len = len(str(cell.value)) if cell.value else 0
                max_len = max(max_len, val_len)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, max_len + 2))


def run_stage5(settings: dict) -> None:
    validated_dir  = settings["paths"]["validated"]
    run_log_path   = settings["paths"]["run_log"]
    manifest_path  = settings["paths"]["manifest"]
    output_xlsx    = settings["paths"]["results_master"]
    output_csv     = output_xlsx.replace(".xlsx", ".csv")

    logger.info("=" * 60)
    logger.info("STAGE 5: Export")
    logger.info("=" * 60)

    manifest      = read_manifest(manifest_path)
    results       = _load_validated_results(validated_dir)
    run_log       = _load_run_log(run_log_path)
    model_names   = list(settings["models"].keys())
    display_names = _model_display_names(settings)

    if not results:
        logger.warning("No validated results found. Run Stages 3 and 4 first.")
        return

    present_models = sorted({model for (model, _) in results.keys()})
    logger.info(f"Models with results: {present_models}")
    logger.info(f"Images in manifest:  {len(manifest)}")

    wb = openpyxl.Workbook()

    # Sheet 1: Main Results
    ws1 = wb.active
    ws1.title = "Main Results"
    headers = _build_sheet1_headers(model_names, display_names)
    ws1.append(headers)
    _apply_header_style(ws1, 1, HEADER_FILL)
    ws1.freeze_panes = "A2"

    sheet1_rows = []
    for mrow in manifest:
        row_data   = _build_sheet1_row(mrow, results, model_names, display_names)
        row_values = [row_data.get(h, "") for h in headers]
        ws1.append(row_values)
        sheet1_rows.append(row_data)

    _auto_width(ws1)

    # Sheet 2: Audit Log
    ws2 = wb.create_sheet("Audit Log")
    if run_log:
        audit_headers = list(run_log[0].keys())
        ws2.append(audit_headers)
        _apply_header_style(ws2, 1, HEADER_FILL)
        for log_row in run_log:
            ws2.append([log_row.get(h, "") for h in audit_headers])
    else:
        ws2.append(["No run log data available"])
    _auto_width(ws2)

    Path(output_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)
    logger.info(f"Saved: {output_xlsx}")

    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(sheet1_rows)
    logger.info(f"Saved: {output_csv}")

    logger.info("")
    logger.info("=" * 60)
    logger.info("STAGE 5 COMPLETE")
    logger.info(f"Rows in Main Results: {len(sheet1_rows)}")
    logger.info(f"Rows in Audit Log:    {len(run_log)}")
    logger.info("=" * 60)
